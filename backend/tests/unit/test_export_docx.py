"""Unit tests for app.export.docx."""

from __future__ import annotations

import io

from docx import Document

from app.export.docx import render_resume_docx
from app.resume.models import Resume
from tests.fixtures import make_resume


def test_render_resume_docx_produces_nonempty_valid_docx() -> None:
    data = render_resume_docx(make_resume())
    assert isinstance(data, bytes)
    assert len(data) > 0
    # A real docx: python-docx can open it back up without error.
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Jordan Ellery Vance" in text
    assert "Cascade Systems" in text


def test_render_resume_docx_includes_bullets_and_skills() -> None:
    data = render_resume_docx(make_resume())
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Migrated the billing pipeline to event sourcing" in text
    assert "Python" in text


def test_render_resume_docx_handles_empty_resume() -> None:
    data = render_resume_docx(Resume())
    document = Document(io.BytesIO(data))
    assert len(document.paragraphs) >= 1
