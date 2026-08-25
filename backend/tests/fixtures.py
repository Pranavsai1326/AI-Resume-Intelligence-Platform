"""Synthetic document builders for tests.

Every name, email and employer below is fabricated for testing and does not belong to a real
person. Building files programmatically (rather than committing binary fixtures) keeps the test
suite's inputs auditable in plain Python.
"""

from __future__ import annotations

import io
import zipfile

from app.resume.models import (
    ContactInfo,
    ExperienceEntry,
    Resume,
    SkillGroup,
)
from app.resume.provenance import Provenance, ProvenancedValue

SYNTHETIC_RESUME_TEXT = """Jordan Ellery Vance
jordan.vance@example-fixture.test | (415) 555-0139
San Francisco, CA
linkedin.com/in/jordanvance

SUMMARY
Backend engineer with 6 years building distributed systems.

EXPERIENCE
Senior Backend Engineer, Cascade Systems
Jan 2021 - Present
- Led migration of the billing pipeline to event sourcing
- Reduced p99 latency by tuning the query planner

Software Engineer | Northlight Data
Jun 2018 - Dec 2020
- Built the ingestion service from scratch

EDUCATION
University of Riverbend, B.S. Computer Science
2014 - 2018

SKILLS
Languages: Python, Go, TypeScript
Infrastructure: Kubernetes, Terraform, AWS

PROJECTS
Trailmark (Python, FastAPI)
- Open source resume parser used by 200 developers

CERTIFICATIONS
AWS Certified Solutions Architect - Amazon - 2022

AWARDS
- Hackathon winner 2019
"""


def make_txt_bytes(text: str = SYNTHETIC_RESUME_TEXT) -> bytes:
    return text.encode("utf-8")


def make_pdf_bytes(lines: list[str] | None = None, *, pages: int = 1) -> bytes:
    """A real, parseable single- or multi-page PDF with the given lines of text."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    if lines is None:
        lines = SYNTHETIC_RESUME_TEXT.split("\n")

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=letter)
    _width, height = letter
    for _ in range(pages):
        y = height - 72
        for line in lines:
            doc.drawString(72, y, line)
            y -= 14
        doc.showPage()
    doc.save()
    return buffer.getvalue()


def make_two_column_pdf_bytes() -> bytes:
    """A page with two genuinely separated blocks of text, left and right."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    left_lines = [f"Left column line {i}" for i in range(1, 16)]
    right_lines = [f"Right column line {i}" for i in range(1, 16)]
    y = height - 72
    for left, right in zip(left_lines, right_lines, strict=True):
        doc.drawString(72, y, left)
        doc.drawString(width - 220, y, right)
        y -= 14
    doc.showPage()
    doc.save()
    return buffer.getvalue()


def make_encrypted_pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.pdfencrypt import StandardEncryption
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=letter)
    doc.setEncrypt(StandardEncryption("secret", ownerPassword="secret-owner"))
    doc.drawString(72, 700, "Protected content")
    doc.showPage()
    doc.save()
    return buffer.getvalue()


def make_docx_bytes(
    paragraphs: list[str] | None = None,
    *,
    with_table: bool = True,
    with_two_columns: bool = False,
) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if paragraphs is None:
        paragraphs = SYNTHETIC_RESUME_TEXT.split("\n")

    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)

    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Skill"
        table.rows[0].cells[1].text = "Level"
        table.rows[1].cells[0].text = "Python"
        table.rows[1].cells[1].text = "Expert"

    if with_two_columns:
        section = document.sections[0]
        sect_pr = section._sectPr
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), "2")
        sect_pr.append(cols)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_header_footer_bytes(text: str = "Jordan Ellery Vance") -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Some content.")
    section = document.sections[0]
    section.header.paragraphs[0].text = text
    section.footer.paragraphs[0].text = "Page 1"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_corrupt_docx_bytes() -> bytes:
    """A valid zip that is missing the parts required of a real Word document."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("not_a_docx.txt", "hello")
    return buffer.getvalue()


def make_zip_bomb_docx_bytes(*, entries: int = 5, entry_size: int = 5_000_000) -> bytes:
    """A zip within DOCX-shaped bounds check territory: many/large entries, highly compressible."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document/>")
        for index in range(entries):
            archive.writestr(f"word/bomb_{index}.xml", "A" * entry_size)
    return buffer.getvalue()


def make_random_binary_bytes(size: int = 256) -> bytes:
    import os

    return os.urandom(size)


def make_resume(
    *, summary: str = "Backend engineer with 6 years building distributed systems."
) -> Resume:
    """A small structured resume for unit tests that need a ``Resume`` directly (no extraction)."""
    extracted = Provenance.extracted(confidence=0.9)
    return Resume(
        contact=ContactInfo(
            full_name=ProvenancedValue(value="Jordan Ellery Vance", provenance=extracted),
            email=ProvenancedValue(value="jordan.vance@example-fixture.test", provenance=extracted),
        ),
        summary=ProvenancedValue(value=summary, provenance=extracted) if summary else None,
        experience=[
            ProvenancedValue(
                value=ExperienceEntry(
                    title="Senior Backend Engineer",
                    organization="Cascade Systems",
                    bullets=[
                        "Migrated the billing pipeline to event sourcing",
                        "Reduced latency by tuning the query planner",
                    ],
                ),
                provenance=extracted,
            )
        ],
        skills=[
            ProvenancedValue(
                value=SkillGroup(category="Languages", skills=["Python", "Go", "TypeScript"]),
                provenance=extracted,
            ),
            ProvenancedValue(
                value=SkillGroup(
                    category="Infrastructure", skills=["Kubernetes", "Terraform", "AWS"]
                ),
                provenance=extracted,
            ),
        ],
    )
