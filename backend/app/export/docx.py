"""DOCX export via python-docx.

Text is set through python-docx's object API (``run.text = ...``, paragraph styles), never by
concatenating raw XML strings, so there is no injection surface to escape against
(SECURITY.md section 6) - the same reason ``app.documents.extract.docx`` trusts python-docx for
the structural half of extraction.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt

from app.resume.models import Resume


def render_resume_docx(resume: Resume) -> bytes:
    document = Document()
    document.styles["Normal"].font.size = Pt(10.5)

    name = resume.contact.full_name.value if resume.contact.full_name else "Resume"
    document.add_heading(name, level=1)

    contact_bits: list[str] = []
    if resume.contact.email:
        contact_bits.append(resume.contact.email.value)
    if resume.contact.phone:
        contact_bits.append(resume.contact.phone.value)
    if resume.contact.location:
        contact_bits.append(resume.contact.location.value)
    contact_bits.extend(link.value for link in resume.contact.links)
    if contact_bits:
        document.add_paragraph(" | ".join(contact_bits))

    if resume.summary:
        document.add_paragraph(resume.summary.value)

    if resume.experience:
        document.add_heading("Experience", level=2)
        for experience_item in resume.experience:
            experience_entry = experience_item.value
            header = document.add_paragraph()
            header.add_run(experience_entry.title).bold = True
            if experience_entry.dates:
                header.add_run(f"  ({experience_entry.dates.raw})")
            document.add_paragraph().add_run(experience_entry.organization).italic = True
            for bullet in experience_entry.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    if resume.education:
        document.add_heading("Education", level=2)
        for education_item in resume.education:
            education_entry = education_item.value
            header = document.add_paragraph()
            header.add_run(education_entry.institution).bold = True
            if education_entry.dates:
                header.add_run(f"  ({education_entry.dates.raw})")
            if education_entry.degree:
                document.add_paragraph(education_entry.degree)
            for detail in education_entry.details:
                document.add_paragraph(detail, style="List Bullet")

    if resume.skills:
        document.add_heading("Skills", level=2)
        for item in resume.skills:
            group = item.value
            text = ", ".join(group.skills)
            if group.category:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{group.category}: ").bold = True
                paragraph.add_run(text)
            else:
                document.add_paragraph(text)

    if resume.projects:
        document.add_heading("Projects", level=2)
        for project_item in resume.projects:
            project_entry = project_item.value
            title = project_entry.name
            if project_entry.technologies:
                title += f" ({', '.join(project_entry.technologies)})"
            document.add_paragraph().add_run(title).bold = True
            if project_entry.description:
                document.add_paragraph(project_entry.description)
            for bullet in project_entry.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    if resume.certifications:
        document.add_heading("Certifications", level=2)
        for certification_item in resume.certifications:
            certification_entry = certification_item.value
            bits = [certification_entry.name]
            if certification_entry.issuer:
                bits.append(certification_entry.issuer)
            if certification_entry.date:
                bits.append(certification_entry.date)
            document.add_paragraph(" — ".join(bits), style="List Bullet")

    for custom_item in resume.custom_sections:
        custom_entry = custom_item.value
        document.add_heading(custom_entry.title, level=2)
        for bullet in custom_entry.bullets:
            document.add_paragraph(bullet, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
